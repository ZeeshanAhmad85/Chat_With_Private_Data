"""
DOCX Document Loader using python-docx.
Extracts paragraphs and structured table content with metadata.
"""

import hashlib
import io
from pathlib import Path
from typing import BinaryIO, List, Union
import docx
from rag_core.models import Document, DocumentMetadata


class DocxLoader:
    """Extracts text, headings, and table cells from DOCX files."""

    def load(self, source: Union[str, Path, BinaryIO], filename: str) -> List[Document]:
        if isinstance(source, (str, Path)):
            with open(source, "rb") as f:
                content = f.read()
        elif hasattr(source, "read"):
            content = source.read()
            if hasattr(source, "seek"):
                source.seek(0)
        else:
            raise ValueError(f"Unsupported source type: {type(source)}")

        if not content:
            raise ValueError(f"File '{filename}' is empty (0 bytes).")

        file_hash = hashlib.sha256(content).hexdigest()
        stream = io.BytesIO(content)

        try:
            doc = docx.Document(stream)
        except Exception as e:
            raise ValueError(f"Corrupted or invalid DOCX format in '{filename}': {e}") from e

        extracted_lines = []

        # Extract paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                extracted_lines.append(text)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Represent table rows cleanly as pipe-separated markdown-like rows
                    extracted_lines.append(" | ".join(row_cells))

        combined_text = "\n\n".join(extracted_lines).strip()

        if not combined_text:
            raise ValueError(f"DOCX '{filename}' contains no readable text or paragraphs.")

        metadata = DocumentMetadata(
            source=filename,
            page=1,
            file_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_size_bytes=len(content),
            file_hash=file_hash,
            extra={"num_paragraphs": len(doc.paragraphs), "num_tables": len(doc.tables)},
        )

        return [Document(text=combined_text, metadata=metadata)]
